from __future__ import annotations

import json
import logging
import os
from contextlib import contextmanager
from datetime import date, datetime, timedelta
from pathlib import Path
from threading import RLock
from typing import Any

import boto3
from botocore.exceptions import ClientError
from boto3.exceptions import S3UploadFailedError
from openpyxl import Workbook, load_workbook

from app.config import settings
from app.enums import DeliveryStatus, PaymentStatus, SubscriptionStatus, UserRole
from app.security import hash_password


logger = logging.getLogger(__name__)

SHEETS: dict[str, list[str]] = {
    "users": ["id", "full_name", "email", "mobile_number", "password_hash", "role", "is_active", "created_at"],
    "customer_profiles": [
        "id",
        "user_id",
        "whatsapp_number",
        "delivery_address",
        "gps_location",
        "age",
        "gender",
        "height_cm",
        "weight_kg",
        "goal_weight_kg",
        "dietary_preferences",
        "allergies",
    ],
    "subscription_plans": [
        "id",
        "name",
        "goal",
        "description",
        "duration_days",
        "delivery_days",
        "fruits_included",
        "optional_addons",
        "price",
        "discount",
        "final_price",
        "is_active",
    ],
    "subscriptions": [
        "id",
        "customer_id",
        "plan_id",
        "status",
        "selected_addons",
        "start_date",
        "end_date",
        "pause_from",
        "pause_to",
        "renewed_from_subscription_id",
        "created_at",
    ],
    "health_logs": ["id", "customer_id", "log_date", "weight_kg", "progress_photo_url", "notes"],
    "deliveries": [
        "id",
        "subscription_id",
        "customer_id",
        "delivery_partner_id",
        "delivery_date",
        "status",
        "route_order",
        "confirmation_photo_url",
        "notes",
    ],
    "payments": [
        "id",
        "subscription_id",
        "customer_id",
        "amount",
        "method",
        "status",
        "transaction_reference",
        "paid_at",
        "created_at",
    ],
    "invoices": ["id", "payment_id", "invoice_number", "invoice_url", "issued_at"],
    "inventory": ["id", "item_name", "category", "unit", "quantity", "low_stock_threshold", "updated_at"],
    "whatsapp_messages": ["id", "customer_id", "template_name", "phone_number", "body", "status", "sent_at", "created_at"],
    "notifications": ["id", "user_id", "channel", "title", "body", "is_read", "created_at"],
    "support_tickets": ["id", "customer_id", "subject", "message", "status", "created_at"],
}

JSON_COLUMNS = {
    "dietary_preferences",
    "allergies",
    "delivery_days",
    "fruits_included",
    "optional_addons",
    "selected_addons",
}

DATE_COLUMNS = {"start_date", "end_date", "pause_from", "pause_to", "log_date", "delivery_date"}
DATETIME_COLUMNS = {"created_at", "paid_at", "issued_at", "sent_at", "updated_at"}

DEFAULT_SUBSCRIPTION_PLANS = [
    {
        "name": "Medium Bowl",
        "goal": "Balanced Nutrition",
        "description": "A balanced monthly fruit bowl plan with seasonal fruits, veggies, sprouts, and a boiled egg for everyday wellness.",
        "duration_days": 30,
        "delivery_days": ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"],
        "fruits_included": ["3 Fruits", "2 Veggies", "Sprouts", "1 Boiled Egg"],
        "optional_addons": ["Sprouts", "Dry Fruits", "Chia Seeds", "Protein Mix"],
        "price": 2500,
        "discount": 500,
        "final_price": 2000,
        "is_active": True,
    },
    {
        "name": "Weight Loss Bowl",
        "goal": "Weight Loss",
        "description": "Low-calorie bowls with fiber-rich fruits, sprouts, cucumber, and chia seeds.",
        "duration_days": 30,
        "delivery_days": ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"],
        "fruits_included": ["Apple", "Papaya", "Watermelon", "Guava", "Cucumber", "Sprouts", "Chia Seeds"],
        "optional_addons": ["Boiled Eggs", "Sprouts", "Dry Fruits", "Protein Mix"],
        "price": 3499,
        "discount": 500,
        "final_price": 2999,
        "is_active": True,
    },
    {
        "name": "Weight Gain Bowl",
        "goal": "Weight Gain",
        "description": "Calorie-dense bowls with banana, mango, dates, avocado, nuts, eggs, and protein-rich add-ons.",
        "duration_days": 30,
        "delivery_days": ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"],
        "fruits_included": ["Banana", "Mango", "Dates", "Avocado", "Dry Fruits", "Boiled Eggs", "Peanut Butter"],
        "optional_addons": ["Boiled Eggs", "Sprouts", "Dry Fruits", "Protein Mix"],
        "price": 4299,
        "discount": 600,
        "final_price": 3699,
        "is_active": True,
    },
]


class ExcelStore:
    def __init__(self, path: str | Path):
        self.path = Path(path)
        self.lock = RLock()
        self._s3_client: Any | None = None

    def setup(self) -> None:
        with self.lock, self._file_lock():
            logger.info("Setting up Excel store", extra={"path": str(self.path), "s3_uri": self.s3_uri()})
            downloaded_from_s3 = self._download_from_s3()
            self.path.parent.mkdir(parents=True, exist_ok=True)
            workbook_exists = self.path.exists()
            if workbook_exists:
                workbook = load_workbook(self.path)
            else:
                workbook = Workbook()
                workbook.remove(workbook.active)
            changed = False
            for sheet_name, headers in SHEETS.items():
                if sheet_name not in workbook.sheetnames:
                    worksheet = workbook.create_sheet(sheet_name)
                    worksheet.append(headers)
                    changed = True
                else:
                    worksheet = workbook[sheet_name]
                    if worksheet.max_row == 0:
                        worksheet.append(headers)
                        changed = True
                    else:
                        existing_headers = [cell.value for cell in worksheet[1]]
                        for header in headers:
                            if header not in existing_headers:
                                worksheet.cell(row=1, column=len(existing_headers) + 1).value = header
                                existing_headers.append(header)
                                changed = True
            changed = self._ensure_default_plans(workbook) or changed
            if changed or not workbook_exists:
                workbook.save(self.path)
            if changed or not workbook_exists or (self._s3_enabled() and not downloaded_from_s3):
                self._upload_to_s3()
            logger.info("Excel store ready", extra={"changed": changed, "downloaded_from_s3": downloaded_from_s3})

    def seed(self, force: bool = False) -> None:
        if settings.is_production:
            raise RuntimeError("Demo seed data is disabled in production")
        logger.info("Seeding Excel store", extra={"force": force, "s3_uri": self.s3_uri()})
        if force:
            with self.lock, self._file_lock():
                workbook = Workbook()
                workbook.remove(workbook.active)
                for sheet_name, headers in SHEETS.items():
                    worksheet = workbook.create_sheet(sheet_name)
                    worksheet.append(headers)
                workbook.save(self.path)
                self._upload_to_s3()
        else:
            self.setup()

        if self.find_one("users", email="admin@biteafruit.com"):
            logger.info("Seed skipped because admin user already exists")
            return

        now = datetime.utcnow()
        admin = self.insert(
            "users",
            {
                "full_name": "Bite Admin",
                "email": "admin@biteafruit.com",
                "mobile_number": "+919900000001",
                "password_hash": hash_password("password123"),
                "role": UserRole.admin.value,
                "is_active": True,
                "created_at": now,
            },
        )
        customer = self.insert(
            "users",
            {
                "full_name": "Aarav Sharma",
                "email": "customer@biteafruit.com",
                "mobile_number": "+919900000002",
                "password_hash": hash_password("password123"),
                "role": UserRole.customer.value,
                "is_active": True,
                "created_at": now,
            },
        )
        partner = self.insert(
            "users",
            {
                "full_name": "Nisha Delivery",
                "email": "delivery@biteafruit.com",
                "mobile_number": "+919900000003",
                "password_hash": hash_password("password123"),
                "role": UserRole.delivery_partner.value,
                "is_active": True,
                "created_at": now,
            },
        )
        _ = admin
        self.insert(
            "customer_profiles",
            {
                "user_id": customer["id"],
                "whatsapp_number": "+919900000002",
                "delivery_address": "12 Green Avenue, Bengaluru",
                "gps_location": "12.9716,77.5946",
                "age": 29,
                "gender": "Male",
                "height_cm": 174,
                "weight_kg": 72,
                "goal_weight_kg": 66,
                "dietary_preferences": ["Vegetarian", "Low sugar"],
                "allergies": ["Kiwi"],
            },
        )
        loss_plan = self.find_one("subscription_plans", name="Weight Loss Bowl")
        for item_name, category, unit, quantity, threshold in [
            ("Apple", "Fruit", "kg", 42, 10),
            ("Papaya", "Fruit", "kg", 26, 8),
            ("Banana", "Fruit", "dozen", 18, 10),
            ("Sprouts", "Protein", "kg", 4, 5),
            ("Eggs", "Add-on", "piece", 90, 50),
            ("Dry Fruits", "Add-on", "kg", 3, 4),
            ("Protein Mix", "Add-on", "kg", 7, 3),
        ]:
            self.insert(
                "inventory",
                {
                    "item_name": item_name,
                    "category": category,
                    "unit": unit,
                    "quantity": quantity,
                    "low_stock_threshold": threshold,
                    "updated_at": now,
                },
            )
        start = date.today()
        subscription = self.insert(
            "subscriptions",
            {
                "customer_id": customer["id"],
                "plan_id": loss_plan["id"],
                "status": SubscriptionStatus.active.value,
                "selected_addons": ["Sprouts", "Protein Mix"],
                "start_date": start,
                "end_date": start + timedelta(days=29),
                "created_at": now,
            },
        )
        self.create_weekday_deliveries(subscription, partner["id"])
        self.insert(
            "payments",
            {
                "subscription_id": subscription["id"],
                "customer_id": customer["id"],
                "amount": loss_plan["final_price"],
                "method": "UPI",
                "status": PaymentStatus.paid.value,
                "transaction_reference": "UPI-DEMO-001",
                "paid_at": now,
                "created_at": now,
            },
        )
        for index, weight in enumerate([72, 70.8, 69.9]):
            self.insert(
                "health_logs",
                {
                    "customer_id": customer["id"],
                    "log_date": date.today() - timedelta(days=(14 - index * 7)),
                    "weight_kg": weight,
                    "notes": "Weekly check-in",
                },
            )
        logger.info("Seed data inserted")

    def _ensure_default_plans(self, workbook: Workbook) -> bool:
        worksheet = workbook["subscription_plans"]
        headers = [cell.value for cell in worksheet[1]]
        id_index = headers.index("id") + 1
        name_index = headers.index("name") + 1
        existing_by_name = {
            worksheet.cell(row=row_number, column=name_index).value: row_number
            for row_number in range(2, worksheet.max_row + 1)
            if worksheet.cell(row=row_number, column=name_index).value
        }
        next_id = self._next_id(worksheet, headers)
        changed = False

        for plan in DEFAULT_SUBSCRIPTION_PLANS:
            row_number = existing_by_name.get(plan["name"])
            if row_number:
                for key, value in plan.items():
                    if key in headers:
                        cell = worksheet.cell(row=row_number, column=headers.index(key) + 1)
                        serialized = self._serialize(key, value)
                        if cell.value != serialized:
                            cell.value = serialized
                            changed = True
            else:
                row = {header: None for header in headers}
                row.update(plan)
                row["id"] = next_id
                worksheet.append([self._serialize(header, row.get(header)) for header in headers])
                next_id += 1
                changed = True

        return changed

    def all(self, sheet_name: str) -> list[dict[str, Any]]:
        with self.lock, self._file_lock():
            self._download_from_s3()
            workbook = load_workbook(self.path)
            worksheet = workbook[sheet_name]
            headers = [cell.value for cell in worksheet[1]]
            rows = []
            for row in worksheet.iter_rows(min_row=2, values_only=True):
                if not any(value is not None for value in row):
                    continue
                rows.append({header: self._deserialize(header, value) for header, value in zip(headers, row)})
            return rows

    def get(self, sheet_name: str, row_id: int) -> dict[str, Any] | None:
        return self.find_one(sheet_name, id=row_id)

    def find_one(self, sheet_name: str, **filters: Any) -> dict[str, Any] | None:
        for row in self.all(sheet_name):
            if self._matches(row, filters):
                return row
        return None

    def filter(self, sheet_name: str, **filters: Any) -> list[dict[str, Any]]:
        return [row for row in self.all(sheet_name) if self._matches(row, filters)]

    def insert(self, sheet_name: str, values: dict[str, Any]) -> dict[str, Any]:
        with self.lock, self._file_lock():
            self._download_from_s3()
            workbook = load_workbook(self.path)
            worksheet = workbook[sheet_name]
            headers = [cell.value for cell in worksheet[1]]
            next_id = self._next_id(worksheet, headers)
            row = {header: None for header in headers}
            row.update(values)
            row["id"] = next_id
            worksheet.append([self._serialize(header, row.get(header)) for header in headers])
            workbook.save(self.path)
            self._upload_to_s3()
            logger.info("Inserted Excel row", extra={"sheet": sheet_name, "row_id": next_id})
        return self.get(sheet_name, next_id) or row

    def update(self, sheet_name: str, row_id: int, values: dict[str, Any]) -> dict[str, Any] | None:
        with self.lock, self._file_lock():
            self._download_from_s3()
            workbook = load_workbook(self.path)
            worksheet = workbook[sheet_name]
            headers = [cell.value for cell in worksheet[1]]
            id_index = headers.index("id") + 1
            for row_number in range(2, worksheet.max_row + 1):
                if worksheet.cell(row=row_number, column=id_index).value == row_id:
                    for key, value in values.items():
                        if key in headers:
                            worksheet.cell(row=row_number, column=headers.index(key) + 1).value = self._serialize(key, value)
                    workbook.save(self.path)
                    self._upload_to_s3()
                    logger.info("Updated Excel row", extra={"sheet": sheet_name, "row_id": row_id})
                    return {
                        header: self._deserialize(header, worksheet.cell(row=row_number, column=index + 1).value)
                        for index, header in enumerate(headers)
                    }
        logger.warning("Excel row update target not found", extra={"sheet": sheet_name, "row_id": row_id})
        return None

    def create_weekday_deliveries(self, subscription: dict[str, Any], partner_id: int | None = None) -> None:
        with self.lock, self._file_lock():
            self._download_from_s3()
            workbook = load_workbook(self.path)
            worksheet = workbook["deliveries"]
            headers = [cell.value for cell in worksheet[1]]
            next_id = self._next_id(worksheet, headers)
            cursor = subscription["start_date"]
            end_date = subscription["end_date"]
            route_order = 1
            created = 0

            while cursor <= end_date:
                if cursor.weekday() < 5:
                    row = {header: None for header in headers}
                    row.update(
                        {
                            "id": next_id,
                            "subscription_id": subscription["id"],
                            "customer_id": subscription["customer_id"],
                            "delivery_partner_id": partner_id,
                            "delivery_date": cursor,
                            "status": DeliveryStatus.pending.value,
                            "route_order": route_order,
                        }
                    )
                    worksheet.append([self._serialize(header, row.get(header)) for header in headers])
                    next_id += 1
                    route_order += 1
                    created += 1
                cursor += timedelta(days=1)

            workbook.save(self.path)
            self._upload_to_s3()
            logger.info(
                "Created weekday deliveries",
                extra={"subscription_id": subscription["id"], "delivery_count": created},
            )

    def _next_id(self, worksheet: Any, headers: list[str]) -> int:
        id_index = headers.index("id") + 1
        ids = [worksheet.cell(row=row_number, column=id_index).value for row_number in range(2, worksheet.max_row + 1)]
        numeric_ids = [int(value) for value in ids if value is not None]
        return max(numeric_ids, default=0) + 1

    def _matches(self, row: dict[str, Any], filters: dict[str, Any]) -> bool:
        return all(row.get(key) == value for key, value in filters.items())

    def _serialize(self, header: str, value: Any) -> Any:
        if value is None:
            return None
        if header in JSON_COLUMNS:
            return json.dumps(value)
        if isinstance(value, (date, datetime)):
            return value.isoformat()
        if hasattr(value, "value"):
            return value.value
        return value

    def _deserialize(self, header: str, value: Any) -> Any:
        if value is None:
            return None
        if header in JSON_COLUMNS:
            if isinstance(value, str):
                return json.loads(value)
            return value
        if header in DATE_COLUMNS:
            if isinstance(value, datetime):
                return value.date()
            if isinstance(value, date):
                return value
            return date.fromisoformat(str(value))
        if header in DATETIME_COLUMNS:
            if isinstance(value, datetime):
                return value
            if isinstance(value, date):
                return datetime.combine(value, datetime.min.time())
            return datetime.fromisoformat(str(value))
        return value

    def _s3_enabled(self) -> bool:
        return settings.use_excel_s3

    def s3_enabled(self) -> bool:
        return self._s3_enabled()

    def s3_uri(self) -> str | None:
        if not self._s3_enabled():
            return None
        return f"s3://{settings.excel_s3_bucket}/{settings.excel_s3_key}"

    def sync_to_s3(self) -> None:
        self._upload_to_s3()

    @contextmanager
    def _file_lock(self):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        lock_path = self.path.with_suffix(f"{self.path.suffix}.lock")
        with lock_path.open("a+b") as lock_file:
            lock_file.seek(0)
            if not lock_file.read(1):
                lock_file.write(b"\0")
                lock_file.flush()
            lock_file.seek(0)

            if os.name == "nt":
                import msvcrt

                msvcrt.locking(lock_file.fileno(), msvcrt.LK_LOCK, 1)
                try:
                    yield
                finally:
                    lock_file.seek(0)
                    msvcrt.locking(lock_file.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                import fcntl

                fcntl.flock(lock_file.fileno(), fcntl.LOCK_EX)
                try:
                    yield
                finally:
                    fcntl.flock(lock_file.fileno(), fcntl.LOCK_UN)

    def _get_s3_client(self) -> Any:
        if self._s3_client is None:
            kwargs = {"region_name": settings.aws_region} if settings.aws_region else {}
            self._s3_client = boto3.client("s3", **kwargs)
        return self._s3_client

    def _s3_permission_error(self, operation: str) -> RuntimeError:
        return RuntimeError(
            f"AWS S3 {operation} was forbidden for {self.s3_uri()}. "
            "Check that the bucket belongs to your AWS account or is accessible by it, "
            "the object key is correct, and the active AWS identity has s3:GetObject "
            "and s3:PutObject permission for this workbook."
        )

    def _download_from_s3(self) -> bool:
        if not self._s3_enabled():
            return False
        self.path.parent.mkdir(parents=True, exist_ok=True)
        try:
            logger.info("Downloading Excel workbook from S3", extra={"s3_uri": self.s3_uri(), "path": str(self.path)})
            self._get_s3_client().download_file(settings.excel_s3_bucket, settings.excel_s3_key, str(self.path))
            logger.info("Downloaded Excel workbook from S3", extra={"s3_uri": self.s3_uri()})
            return True
        except ClientError as exc:
            error_code = exc.response.get("Error", {}).get("Code")
            if error_code in {"404", "NoSuchKey", "NotFound"}:
                logger.info("Excel workbook does not exist in S3 yet", extra={"s3_uri": self.s3_uri()})
                return False
            if error_code in {"403", "AccessDenied"}:
                logger.warning("S3 download forbidden", extra={"s3_uri": self.s3_uri()})
                raise self._s3_permission_error("download") from exc
            raise

    def _upload_to_s3(self) -> None:
        if not self._s3_enabled():
            return
        try:
            logger.info("Uploading Excel workbook to S3", extra={"s3_uri": self.s3_uri(), "path": str(self.path)})
            self._get_s3_client().upload_file(
                str(self.path),
                settings.excel_s3_bucket,
                settings.excel_s3_key,
                ExtraArgs={"ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
            )
            logger.info("Uploaded Excel workbook to S3", extra={"s3_uri": self.s3_uri()})
        except ClientError as exc:
            error_code = exc.response.get("Error", {}).get("Code")
            if error_code in {"403", "AccessDenied"}:
                logger.warning("S3 upload forbidden", extra={"s3_uri": self.s3_uri()})
                raise self._s3_permission_error("upload") from exc
            raise
        except S3UploadFailedError as exc:
            logger.warning("S3 upload failed", extra={"s3_uri": self.s3_uri()})
            raise self._s3_permission_error("upload") from exc


store = ExcelStore(settings.excel_database_cache_path)
